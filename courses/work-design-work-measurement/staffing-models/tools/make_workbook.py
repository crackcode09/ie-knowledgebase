"""make_workbook.py - build templates/staffing-model-workbook.xlsx (Excel version of the course model).

    python tools/make_workbook.py

Yellow cells are inputs. All results are live formulas.
License: MIT
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, os.path.dirname(__file__))
import stmkit as sk  # noqa: E402

OUT = os.path.join(os.path.dirname(__file__), "..", "templates", "staffing-model-workbook.xlsx")
B = Font(bold=True)
IN = PatternFill("solid", fgColor="FFF2CC")
HEAD = PatternFill("solid", fgColor="DDEBF7")


def head(ws, row, labels):
    for j, lab in enumerate(labels, 1):
        c = ws.cell(row=row, column=j, value=lab)
        c.font, c.fill = B, HEAD


def inp(ws, ref, v):
    ws[ref] = v
    ws[ref].fill = IN


def main():
    wb = Workbook()
    av = sk.ex_availability()["average"]
    u = sk.ut()
    v = sk.forecast_all("base")[2026]
    _, _, us = sk.inspector_unit_share()

    # ---- Availability
    a = wb.active
    a.title = "Availability"
    a["A1"], a["A1"].font = "Available hours per FTE (average of 2025 leave records)", B
    rows = [("Paid hours", av["paid_hours"]), ("Holidays", av["holidays"]), ("Annual leave", av["annual_leave"]),
            ("Sick leave", av["sick_leave"]), ("Training", av["training"]), ("Other leave", av["other_leave"])]
    for i, (lab, val) in enumerate(rows, 3):
        a.cell(row=i, column=1, value=lab)
        inp(a, f"B{i}", float(val))
    a["A10"], a["B10"] = "Available hours per FTE", "=B3-SUM(B4:B8)"
    a["A11"], a["B11"] = "Availability %", "=B10/B3"
    a["A13"], a["A13"].font = "Relief factors", B
    a["A14"] = "Hours a post must be covered per year"
    inp(a, "B14", 2000)
    a["A15"], a["B15"] = "FTE per post (relief factor)", "=B14/B10"
    a["A16"], a["B16"] = "FTE per 24/7 post", "=8760/B10"
    a.column_dimensions["A"].width = 40

    # ---- Inspector model
    m = wb.create_sheet("InspectorModel")
    m["A1"], m["A1"].font = "Inspector workload model (2026 base forecast)", B
    head(m, 3, ["Driver", "Annual volume", "Unit time (min)", "Travel per stop (min)", "Hours"])
    lines = [("Residential inspections", v["res_inspections"], u["inspection_res"]),
             ("Commercial inspections", v["com_inspections"], u["inspection_com"]),
             ("Re-inspections", v["reinspections"], u["reinspection"]),
             ("Complaints", v["complaints"], u["complaint"])]
    for i, (lab, vol, t) in enumerate(lines, 4):
        m.cell(row=i, column=1, value=lab)
        inp(m, f"B{i}", round(float(vol)))
        inp(m, f"C{i}", float(t))
        m.cell(row=i, column=4, value="=$B$11")
        m.cell(row=i, column=5, value=f"=B{i}*(C{i}+D{i})/60")
    m["A8"], m["E8"] = "Unit-driven hours", "=SUM(E4:E7)"
    m["A10"], m["A10"].font = "Inputs", B
    m["A11"] = "Travel per stop (min)"
    inp(m, "B11", float(u["travel_per_stop"]))
    m["A12"] = "Unit-driven share of staffed time (work sampling)"
    inp(m, "B12", round(float(us), 3))
    m["A14"], m["B14"] = "Total hours incl. indirect", "=E8/B12"
    m["A15"], m["B15"] = "Available hours per FTE", "=Availability!B10"
    m["A16"], m["B16"] = "Required inspector FTE", "=B14/B15"
    m["A17"] = "On staff"
    inp(m, "B17", 21)
    m["A18"], m["B18"] = "Gap (on staff - required)", "=B17-B16"
    m.column_dimensions["A"].width = 46
    for c in "BCDE":
        m.column_dimensions[c].width = 16

    # ---- Reviewer model
    r = wb.create_sheet("ReviewerModel")
    r["A1"], r["A1"].font = "Plan reviewer workload model (2026 base forecast)", B
    head(r, 3, ["Driver", "Annual volume", "Review (min)", "Resubmittal rate", "Resubmittal (min)", "Hours"])
    lines = [("Residential plan reviews", v["res_plan_reviews"], u["plan_review_res"], sk.RESUB["res"], u["resubmittal_res"]),
             ("Commercial plan reviews", v["com_plan_reviews"], u["plan_review_com"], sk.RESUB["com"], u["resubmittal_com"])]
    for i, (lab, vol, t, rr, rt) in enumerate(lines, 4):
        r.cell(row=i, column=1, value=lab)
        inp(r, f"B{i}", round(float(vol)))
        inp(r, f"C{i}", float(t))
        inp(r, f"D{i}", rr)
        inp(r, f"E{i}", float(rt))
        r.cell(row=i, column=6, value=f"=B{i}*(C{i}+D{i}*E{i})/60")
    r["A7"] = "Indirect share"
    inp(r, "B7", sk.INDIRECT["Plan reviewer"])
    r["A8"], r["B8"] = "Total hours", "=SUM(F4:F5)/(1-B7)"
    r["A9"], r["B9"] = "Required reviewer FTE", "=B8/Availability!B10"
    r.column_dimensions["A"].width = 30

    # ---- Hiring
    h = wb.create_sheet("Hiring")
    h["A1"], h["A1"].font = "Inspector supply projection (yellow = inputs)", B
    h["A2"] = "Quit rate"
    inp(h, "B2", sk.separation_rates()["Inspector"]["quit_rate"])
    h["A3"] = "New-hire first-year productivity"
    inp(h, "B3", sk.NEW_HIRE_FACTOR)
    head(h, 5, ["Year", "Start headcount", "Retirements", "Quits", "Hires", "End headcount", "Effective FTE",
                "Required FTE", "Gap"])
    proj = sk.project_inspectors(tuple(sk.ex_gap()["plan_capped_4"]))
    need = sk.ex_gap()["need"]
    for i, p in enumerate(proj, 6):
        h.cell(row=i, column=1, value=p["year"])
        if i == 6:
            inp(h, "B6", 21)
        else:
            h.cell(row=i, column=2, value=f"=F{i - 1}")
        inp(h, f"C{i}", p["retirements"])
        h.cell(row=i, column=4, value=f"=B{i}*$B$2")
        inp(h, f"E{i}", p["hires"])
        h.cell(row=i, column=6, value=f"=B{i}-C{i}-D{i}+E{i}")
        h.cell(row=i, column=7, value=f"=(B{i}+B{i}-C{i}-D{i})/2+E{i}*$B$3")
        inp(h, f"H{i}", round(need[p["year"]], 2))
        h.cell(row=i, column=9, value=f"=G{i}-H{i}")
    h["A10"] = "Retirements are expected values from the eligibility rule in tools/stmkit.py (illustrative rule)."
    h.column_dimensions["A"].width = 32

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print("wrote", os.path.normpath(OUT))


if __name__ == "__main__":
    main()
