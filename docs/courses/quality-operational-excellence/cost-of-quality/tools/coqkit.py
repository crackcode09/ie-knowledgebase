"""coqkit.py - cost-of-quality data, calculations and tracker for the Cost of Quality course.

Running case: Ellison Hydraulics (fictional) builds hydraulic cylinders; about $40M annual sales.
A cost-of-quality (CoQ) program starts in month 7 of the year.

    python tools/coqkit.py data        # (re)write data/*.csv (fixed seed)
    python tools/coqkit.py all         # every worked example
    python tools/coqkit.py <name>      # summary, monthly, ratios, failures, drivers, hidden, loss, projects, model
    python tools/coqkit.py tracker     # build templates/coq-tracker.xlsx (Excel tracker with live formulas)
    python tools/coqkit.py report      # write solutions/coq-summary-report.md from the data

Requires numpy, pandas; openpyxl for the Excel tracker. License: MIT
"""
from __future__ import annotations

import os
import sys

import numpy as np
import pandas as pd

ROOT = os.path.join(os.path.dirname(__file__), "..")
D = os.path.join(ROOT, "data")
CATS = ["Prevention", "Appraisal", "Internal failure", "External failure"]

# cost element -> (category, monthly base $ before program, monthly change after program starts, noise SD)
ELEMENTS = {
    # prevention
    "Quality planning & engineering": ("Prevention", 9000, +6000, 400),
    "Training in quality methods": ("Prevention", 1500, +3500, 300),
    "Supplier quality development": ("Prevention", 2000, +4000, 300),
    "Design & process FMEA / reviews": ("Prevention", 1200, +3000, 250),
    "Preventive maintenance (quality-critical equipment)": ("Prevention", 4000, +2500, 400),
    # appraisal
    "Incoming inspection": ("Appraisal", 11000, -2500, 600),
    "In-process inspection": ("Appraisal", 18000, -3000, 900),
    "Final test & inspection": ("Appraisal", 22000, -2000, 900),
    "Calibration & gauge maintenance": ("Appraisal", 3500, 0, 250),
    "Quality audits": ("Appraisal", 2500, 0, 250),
    # internal failure
    "Scrap": ("Internal failure", 52000, -22000, 4000),
    "Rework & repair": ("Internal failure", 38000, -15000, 3000),
    "Re-inspection & re-test": ("Internal failure", 9000, -4000, 900),
    "Failure investigation & MRB": ("Internal failure", 6000, -1500, 600),
    "Supplier-reject handling": ("Internal failure", 4000, -1500, 500),
    # external failure
    "Warranty claims": ("External failure", 30000, -9000, 4000),
    "Returns processing & replacement": ("External failure", 14000, -5000, 1500),
    "Field service visits": ("External failure", 9000, -3000, 1200),
    "Complaint handling": ("External failure", 5000, -1500, 500),
    "Customer concessions & credits": ("External failure", 6000, -2500, 1200),
}


# ---------------------------------------------------------------- data
def make_data(seed=1951):
    rng = np.random.default_rng(seed)
    os.makedirs(D, exist_ok=True)
    months = pd.period_range("2026-01", periods=12, freq="M").astype(str)
    rows = []
    for i, m in enumerate(months):
        ramp = 0.0 if i < 6 else min(1.0, (i - 5) / 4)          # program starts month 7, full effect by month 10
        for el, (cat, base, delta, sd) in ELEMENTS.items():
            lag = ramp if cat in ("Prevention", "Appraisal") else max(0.0, (i - 6) / 5) if i >= 6 else 0.0
            cost = base + delta * (ramp if cat == "Prevention" else lag) + rng.normal(0, sd)
            rows.append(dict(month=m, category=cat, element=el, cost=round(max(cost, 0), -1)))
    pd.DataFrame(rows).to_csv(os.path.join(D, "coq_ledger.csv"), index=False)

    sales = np.round(3_300_000 + rng.normal(0, 120_000, 12) + np.arange(12) * 8_000, -3)
    units = np.round(sales / 1_650).astype(int)
    pd.DataFrame(dict(month=months, sales=sales, units_shipped=units, direct_labor_hours=np.round(units * 6.2))).to_csv(
        os.path.join(D, "monthly_operations.csv"), index=False)

    # failure events (first half year) with drivers, for Pareto and driver analysis
    drivers = {"Weld porosity": ("Welding", 0.24, 1400), "Rod seal leak": ("Assembly", 0.20, 900),
               "Bore out of tolerance": ("Machining", 0.17, 2200), "Paint defects": ("Finishing", 0.13, 300),
               "Wrong port orientation": ("Assembly", 0.09, 650), "Supplier rod plating": ("Supplier", 0.08, 1700),
               "Documentation error": ("Order entry", 0.05, 250), "Handling damage": ("Shipping", 0.04, 500)}
    names = list(drivers)
    probs = np.array([drivers[k][1] for k in names])
    ev = []
    for e in range(1, 701):
        k = names[rng.choice(len(names), p=probs / probs.sum())]
        area, p, mean = drivers[k]
        where = "External" if rng.uniform() < (0.22 if k in ("Rod seal leak", "Supplier rod plating", "Handling damage") else 0.08) else "Internal"
        cost = mean * rng.lognormal(0, 0.5) * (6.0 if where == "External" else 1.0)
        ev.append(dict(event=e, month=str(months[rng.integers(0, 6)]), failure_type=k, process_area=area,
                       detected=where, cost=round(cost, -1)))
    ev = pd.DataFrame(ev)
    # reconcile: event-level costs must add up to the ledger's failure cost for months 1-6
    led = pd.DataFrame(rows)
    target = led[led.month.isin(months[:6]) & led.category.isin(["Internal failure", "External failure"])].cost.sum()
    ev["cost"] = (ev.cost * target / ev.cost.sum()).round(-1)
    ev.to_csv(os.path.join(D, "failure_events.csv"), index=False)

    # hidden-cost survey inputs (annual, before program)
    pd.DataFrame([
        ("Expediting and premium freight for remakes", 96_000, "freight invoices flagged 'quality'"),
        ("Lost capacity on bottleneck (hours x margin per hour)", 310_000, "MES downtime coded quality x $1,150/h"),
        ("Engineering time on firefighting", 140_000, "time sampling of 6 engineers over 4 weeks"),
        ("Excess safety stock held because of yield loss", 45_000, "inventory carrying cost on buffer"),
        ("Sales lost from two customers who reduced orders", 420_000, "sales manager estimate (contribution margin)"),
        ("Management time in customer escalation meetings", 38_000, "calendar review"),
    ], columns=["hidden_cost", "annual_estimate", "basis"]).to_csv(os.path.join(D, "hidden_costs.csv"), index=False)
    return "data written"


def load(name):
    return pd.read_csv(os.path.join(D, name))


# ---------------------------------------------------------------- core calculations
def by_category(ledger, months=None):
    l = ledger if months is None else ledger[ledger.month.isin(months)]
    return l.groupby("category").cost.sum().reindex(CATS)


def halves():
    led = load("coq_ledger.csv")
    m = sorted(led.month.unique())
    return led, m[:6], m[6:9], m[9:]


# ---------------------------------------------------------------- examples
def ex_summary():
    led = load("coq_ledger.csv")
    ops = load("monthly_operations.csv")
    tot = by_category(led)
    sales = ops.sales.sum()
    coq = tot.sum()
    return {"annual CoQ by category ($)": tot.round(0).to_dict(), "total CoQ ($)": round(coq),
            "category share %": (100 * tot / coq).round(1).to_dict(), "annual sales ($)": int(sales),
            "CoQ as % of sales": round(100 * coq / sales, 2),
            "cost of good quality (P + A) vs poor quality (IF + EF) ($)": (round(tot[CATS[:2]].sum()), round(tot[CATS[2:]].sum()))}


def ex_monthly():
    led, h1, ramp, h2 = halves()
    ops = load("monthly_operations.csv").set_index("month")
    out = {}
    for label, ms in (("months 1-6 (before)", h1), ("months 10-12 (after)", h2)):
        t = by_category(led, ms) / len(ms)
        s = ops.loc[ms, "sales"].mean()
        out[label] = {**{f"{k} / month": round(v) for k, v in t.items()}, "total / month": round(t.sum()),
                      "CoQ % of sales": round(100 * t.sum() / s, 2)}
    b, a = out["months 1-6 (before)"]["total / month"], out["months 10-12 (after)"]["total / month"]
    out["monthly reduction ($)"] = b - a
    out["annualized saving at the 'after' run rate ($)"] = 12 * (b - a)
    return out


def ex_ratios():
    led = load("coq_ledger.csv")
    ops = load("monthly_operations.csv").set_index("month")
    m = led.pivot_table(index="month", columns="category", values="cost", aggfunc="sum")[CATS]
    m["total"] = m.sum(axis=1)
    r = pd.DataFrame(index=m.index)
    r["CoQ % sales"] = 100 * m.total / ops.sales
    r["CoQ per unit $"] = m.total / ops.units_shipped
    r["CoQ per labor hour $"] = m.total / ops.direct_labor_hours
    r["failure % of CoQ"] = 100 * (m["Internal failure"] + m["External failure"]) / m.total
    r["prevention % of CoQ"] = 100 * m.Prevention / m.total
    return {"monthly bases and ratios": r.round(2).reset_index().to_dict("records")}


def ex_failures():
    led = load("coq_ledger.csv")
    f = led[led.category.isin(CATS[2:])].groupby(["category", "element"]).cost.sum().sort_values(ascending=False)
    tot = f.sum()
    tab = pd.DataFrame(dict(cost=f, pct=(100 * f / tot).round(1), cum_pct=(100 * f.cumsum() / tot).round(1)))
    return {"failure-cost elements (annual)": tab.reset_index().to_dict("records")}


def ex_drivers():
    ev = load("failure_events.csv")
    g = ev.groupby("failure_type").agg(events=("event", "count"), cost=("cost", "sum"),
                                       external_share=("detected", lambda s: (s == "External").mean()))
    g = g.sort_values("cost", ascending=False)
    share = 100 * g.cost / g.cost.sum()
    g["cost %"] = share.round(1)
    g["cum %"] = share.cumsum().round(1)
    g["avg cost per event"] = (g.cost / g.events).round(0)
    g["external_share"] = (100 * g.external_share).round(1)
    by_count = ev.failure_type.value_counts().index[:3].tolist()
    ext = ev.groupby("detected").agg(events=("event", "count"), cost=("cost", "sum"))
    ext["avg"] = (ext.cost / ext.events).round(0)
    return {"events / total cost (6 months)": (len(ev), round(ev.cost.sum())),
            "drivers ranked by cost": g.reset_index().to_dict("records"),
            "top 3 by count (for contrast)": by_count,
            "internal vs external": ext.round(0).reset_index().to_dict("records"),
            "by process area ($)": ev.groupby("process_area").cost.sum().sort_values(ascending=False).round(0).to_dict()}


def ex_hidden():
    h = load("hidden_costs.csv")
    visible_failure = ex_summary()["annual CoQ by category ($)"]
    vf = visible_failure["Internal failure"] + visible_failure["External failure"]
    # visible failure for the same period (annualized from months 1-6)
    led, h1, _, _ = halves()
    vf_h1 = 2 * by_category(led, h1)[CATS[2:]].sum()
    hidden = h.annual_estimate.sum()
    return {"hidden-cost items": h.to_dict("records"), "hidden total ($/yr)": int(hidden),
            "visible failure cost, annualized from months 1-6 ($/yr)": round(vf_h1),
            "hidden-to-visible multiplier": round(hidden / vf_h1, 2),
            "true failure cost ($/yr)": round(vf_h1 + hidden),
            "note": "estimates have wide uncertainty; report them separately from booked costs"}


def ex_loss():
    """Quadratic (Taguchi) loss for cylinder bore: target 50.000 mm, spec +/-0.025, loss at limit = $180 rework."""
    k = 180 / 0.025 ** 2
    out = {"k ($ per mm^2)": round(k)}
    for name, mu, sd in (("centred, sd 0.008", 50.000, 0.008), ("offset +0.010, sd 0.008", 50.010, 0.008),
                         ("centred, sd 0.004", 50.000, 0.004)):
        loss = k * (sd ** 2 + (mu - 50.0) ** 2)
        out[f"expected loss per cylinder: {name}"] = round(loss, 2)
    out["annual at 24,000 cylinders, offset case vs centred-tight case ($)"] = (
        round(24000 * k * (0.008 ** 2 + 0.010 ** 2)), round(24000 * k * 0.004 ** 2))
    return out


def ex_projects():
    """Candidate improvement projects ranked by CoQ return."""
    p = pd.DataFrame([
        ("Weld porosity: gas-flow monitoring + welder certification", "Weld porosity", 62_000, 0.70),
        ("Rod seal leak: install tool + seal supplier change", "Rod seal leak", 38_000, 0.65),
        ("Bore tolerance: in-process gauging on hone", "Bore out of tolerance", 85_000, 0.60),
        ("Paint line: pretreatment control", "Paint defects", 24_000, 0.50),
        ("Port orientation: error-proof fixture", "Wrong port orientation", 9_000, 0.90),
    ], columns=["project", "driver", "one_time_cost", "expected_reduction"])
    ev = load("failure_events.csv")
    annual = 2 * ev.groupby("failure_type").cost.sum()                  # 6 months -> annual
    p["annual_failure_cost"] = p.driver.map(annual).round(0)
    p["annual_saving"] = (p.annual_failure_cost * p.expected_reduction).round(0)
    p["payback_months"] = (12 * p.one_time_cost / p.annual_saving).round(1)
    p["first_year_ROI_%"] = (100 * (p.annual_saving - p.one_time_cost) / p.one_time_cost).round(0)
    p = p.sort_values("payback_months")
    return {"projects (ranked by payback)": p.drop(columns="driver").to_dict("records"),
            "portfolio: cost, annual saving": (int(p.one_time_cost.sum()), int(p.annual_saving.sum()))}


def ex_model():
    """Classic vs modern view of the conformance / nonconformance trade-off (illustrative)."""
    q = np.linspace(0.90, 0.9999, 500)                 # conformance level
    nonconf = 2_000_000 * (1 - q)                      # failure cost falls with conformance
    classic = 20_000 / (1 - q) ** 0.5 - 20_000 / 0.1 ** 0.5   # prevention+appraisal rising steeply near perfection
    modern = 90_000 * (q - 0.90) / 0.0999              # gradual: prevention gets cheaper as processes are designed right
    i_c = int(np.argmin(classic + nonconf))
    i_m = int(np.argmin(modern + nonconf))
    return {"classic model optimum conformance": round(float(q[i_c]), 4), "classic optimum total ($/period)": round(float((classic + nonconf)[i_c]), -2),
            "modern model optimum conformance": round(float(q[i_m]), 4), "modern optimum total ($/period)": round(float((modern + nonconf)[i_m]), -2)}


# ---------------------------------------------------------------- tracker workbook and report
def build_tracker(path=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter
    path = path or os.path.join(ROOT, "templates", "coq-tracker.xlsx")
    led = load("coq_ledger.csv")
    ops = load("monthly_operations.csv")
    months = sorted(led.month.unique())
    wb = Workbook()
    ws = wb.active
    ws.title = "Ledger"
    ws.append(["Category", "Element"] + months + ["Year"])
    bold = Font(bold=True)
    for c in ws[1]:
        c.font = bold
    piv = led.pivot_table(index=["category", "element"], columns="month", values="cost", aggfunc="sum")
    piv = piv.reindex(sorted(piv.index, key=lambda x: (CATS.index(x[0]), x[1])))
    for (cat, el), row in piv.iterrows():
        r = ws.max_row + 1
        ws.append([cat, el] + [float(row[m]) for m in months])
        ws.cell(r, len(months) + 3, f"=SUM(C{r}:{get_column_letter(len(months) + 2)}{r})")
    last = max(ws.max_row, 500)          # leave room for users to add cost elements
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 48
    s = wb.create_sheet("Summary")
    s.append(["Category"] + months + ["Year"])
    for c in s[1]:
        c.font = bold
    for i, cat in enumerate(CATS, start=2):
        s.cell(i, 1, cat)
        for j in range(len(months) + 1):
            col = get_column_letter(j + 3)
            s.cell(i, j + 2, f'=SUMIF(Ledger!$A$2:$A${last},$A{i},Ledger!{col}$2:{col}${last})')
    tr = len(CATS) + 2
    s.cell(tr, 1, "Total CoQ").font = bold
    for j in range(len(months) + 1):
        col = get_column_letter(j + 2)
        s.cell(tr, j + 2, f"=SUM({col}2:{col}{tr - 1})")
    s.cell(tr + 1, 1, "Sales")
    for j, m in enumerate(months):
        s.cell(tr + 1, j + 2, float(ops.set_index("month").loc[m, "sales"]))
    yc = get_column_letter(len(months) + 2)
    s.cell(tr + 1, len(months) + 2, f"=SUM(B{tr + 1}:{get_column_letter(len(months) + 1)}{tr + 1})")
    s.cell(tr + 2, 1, "CoQ % of sales").font = bold
    s.cell(tr + 3, 1, "Failure % of CoQ")
    s.cell(tr + 4, 1, "Prevention % of CoQ")
    for j in range(len(months) + 1):
        col = get_column_letter(j + 2)
        s.cell(tr + 2, j + 2, f"={col}{tr}/{col}{tr + 1}").number_format = "0.00%"
        s.cell(tr + 3, j + 2, f"=({col}4+{col}5)/{col}{tr}").number_format = "0.0%"
        s.cell(tr + 4, j + 2, f"={col}2/{col}{tr}").number_format = "0.0%"
    s.column_dimensions["A"].width = 22
    note = wb.create_sheet("How to use")
    for line in ["Cost-of-quality tracker (Cost of Quality course).",
                 "1. Ledger: one row per cost element; enter monthly $ in the month columns. Keep the Category text exactly one of:",
                 "   Prevention / Appraisal / Internal failure / External failure.",
                 "2. Summary: totals, CoQ % of sales and mix ratios update automatically (formulas).",
                 "3. Enter monthly sales on the Summary sheet (row 'Sales').",
                 "4. Report hidden-cost estimates separately (see module 07); do not mix them into booked costs.",
                 f"Pre-filled with the fictional Ellison Hydraulics data; clear the month columns to start your own."]:
        note.append([line])
    note.column_dimensions["A"].width = 110
    for w in (ws, s):
        w.freeze_panes = "C2" if w is ws else "B2"
    hdr = PatternFill("solid", fgColor="DDEBF7")
    for w in (ws, s):
        for c in w[1]:
            c.fill = hdr
    wb.save(path)
    return path


def write_report(path=None):
    path = path or os.path.join(ROOT, "solutions", "coq-summary-report.md")
    s, mth, d, h, pr = ex_summary(), ex_monthly(), ex_drivers(), ex_hidden(), ex_projects()
    before, after = mth["months 1-6 (before)"], mth["months 10-12 (after)"]
    lines = ["# Cost of Quality Summary Report — Ellison Hydraulics (fictional), 2026", "",
             "*Generated by `python tools/coqkit.py report` from the course data.*", "",
             "## Headline", "",
             f"- Total booked cost of quality: **${s['total CoQ ($)']:,}** = **{s['CoQ as % of sales']}% of sales**.",
             f"- Monthly CoQ fell from **${before['total / month']:,}** (months 1–6) to **${after['total / month']:,}** (months 10–12): "
             f"{before['CoQ % of sales']}% → {after['CoQ % of sales']}% of sales; annualized saving **${mth['annualized saving at the \'after\' run rate ($)']:,}**.",
             f"- Hidden failure costs are estimated at **${h['hidden total ($/yr)']:,}/yr** — {h['hidden-to-visible multiplier']}× the booked failure cost (estimate; wide uncertainty).", "",
             "## Cost by category (year)", "", "| Category | $ | % of CoQ |", "|---|---|---|"]
    for c in CATS:
        lines.append(f"| {c} | {s['annual CoQ by category ($)'][c]:,.0f} | {s['category share %'][c]} |")
    lines += ["", "## Before vs after the program (monthly averages)", "", "| Category | Months 1–6 | Months 10–12 |", "|---|---|---|"]
    for c in CATS:
        lines.append(f"| {c} | {before[f'{c} / month']:,} | {after[f'{c} / month']:,} |")
    lines.append(f"| **Total** | **{before['total / month']:,}** | **{after['total / month']:,}** |")
    lines += ["", "## Top failure-cost drivers (months 1–6)", "", "| Driver | Events | Cost $ | Cum % |", "|---|---|---|---|"]
    for r in d["drivers ranked by cost"][:5]:
        lines.append(f"| {r['failure_type']} | {r['events']} | {r['cost']:,.0f} | {r['cum %']} |")
    lines += ["", "## Recommended projects", "", "| Project | One-time cost | Annual saving | Payback (months) |", "|---|---|---|---|"]
    for r in pr["projects (ranked by payback)"]:
        lines.append(f"| {r['project']} | {r['one_time_cost']:,} | {r['annual_saving']:,.0f} | {r['payback_months']} |")
    lines += ["", "## Notes", "", "- Booked costs come from the ledger (`data/coq_ledger.csv`); hidden costs are estimates and are reported separately.",
              "- Savings in failure categories are partly offset by higher prevention spending; the net figures above include that.", ""]
    with open(path, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
    return path


def _clean(o):
    if isinstance(o, dict):
        return {k: _clean(v) for k, v in o.items()}
    if isinstance(o, (list, tuple)):
        return type(o)(_clean(v) for v in o)
    if isinstance(o, np.generic):
        return o.item()
    return o


EX = {"summary": ex_summary, "monthly": ex_monthly, "ratios": ex_ratios, "failures": ex_failures, "drivers": ex_drivers,
      "hidden": ex_hidden, "loss": ex_loss, "projects": ex_projects, "model": ex_model}

if __name__ == "__main__":
    which = sys.argv[1] if len(sys.argv) > 1 else "all"
    if which == "data":
        print(make_data())
        sys.exit()
    if which == "tracker":
        print("written", build_tracker())
        sys.exit()
    if which == "report":
        print("written", write_report())
        sys.exit()
    for k, fn in EX.items():
        if which in ("all", k):
            print(f"\n=== {k}")
            for a, b in fn().items():
                if isinstance(b, list) and b and isinstance(b[0], dict):
                    print(f"  {a}:")
                    for row in b:
                        print(f"    {_clean(row)}")
                else:
                    print(f"  {a}: {_clean(b)}")
